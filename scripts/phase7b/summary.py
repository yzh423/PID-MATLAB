"""Build the fixed-layout Phase 7B one-page research summary."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from scripts.phase7b.office import normalize_openxml_package


INK = RGBColor(0, 0, 0)
MUTED = RGBColor(82, 89, 96)
BLUE = RGBColor(31, 78, 121)
RESULT_FILL = "E8EEF5"
SOURCE_FILL = "F2F4F7"
CONTENT_WIDTH_DXA = 10512
CELL_MARGIN_DXA = 100


class SummaryBuildError(ValueError):
    """Raised when the resolved presentation package cannot build a summary."""


def _set_run_font(run: Any, size: float, *, bold: bool = False, color: RGBColor = INK) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def _shade(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shade = properties.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        properties.append(shade)
    shade.set(qn("w:fill"), fill)


def _set_cell_margins(cell: Any, *, top: int, start: int, bottom: int, end: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table: Any, widths: Sequence[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise SummaryBuildError("result strip widths must equal the usable page width")
    table.autofit = False
    properties = table._tbl.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")
    table_indent = properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        properties.append(table_indent)
    table_indent.set(qn("w:w"), "0")
    table_indent.set(qn("w:type"), "dxa")
    layout = properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for column, width in zip(grid.gridCol_lst, widths, strict=True):
        column.set(qn("w:w"), str(width))
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            _set_cell_margins(
                cell,
                top=CELL_MARGIN_DXA,
                start=CELL_MARGIN_DXA,
                bottom=CELL_MARGIN_DXA,
                end=CELL_MARGIN_DXA,
            )


def _configure_style(style: Any, *, size: float, before: float, after: float, color: RGBColor = INK, bold: bool = False) -> None:
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.05


def configure_document(document: Document) -> None:
    """Apply standard-business-brief tokens with a named one-page override."""
    if len(document.sections) != 1:
        raise SummaryBuildError("new summary must begin with exactly one section")
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].text = ""
    section.footer.paragraphs[0].text = ""
    _configure_style(document.styles["Normal"], size=10.25, before=0, after=4)
    _configure_style(document.styles["Heading 1"], size=11.25, before=6, after=2, color=BLUE, bold=True)
    _configure_style(document.styles["Caption"], size=8.25, before=2, after=3, color=MUTED)
    document.core_properties.title = "Reliable Robotic Manipulation Through Evidence-Grounded Control"
    document.core_properties.author = "PID vs Fuzzy PID project"
    document.core_properties.subject = "Phase 7B one-page research summary"
    document.core_properties.comments = "Deterministic Phase 7B research-summary pipeline"


def _add_paragraph(document: Document, text: str, *, style: str | None = None, before: float | None = None, after: float | None = None, size: float | None = None, bold: bool = False, color: RGBColor = INK) -> Any:
    paragraph = document.add_paragraph(style=style)
    if before is not None:
        paragraph.paragraph_format.space_before = Pt(before)
    if after is not None:
        paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    _set_run_font(run, size if size is not None else 10.25, bold=bold, color=color)
    return paragraph


def add_masthead(document: Document, summary: Mapping[str, object]) -> None:
    """Create a restrained memo_masthead without the template's bottom rule."""
    label = _add_paragraph(document, "RESEARCH SUMMARY | PHASE 7B", before=0, after=1, size=8.5, bold=True, color=BLUE)
    label.paragraph_format.keep_with_next = True
    title = _add_paragraph(document, str(summary["title"]), before=0, after=2, size=17.5, bold=True)
    title.paragraph_format.keep_with_next = True
    takeaway = _add_paragraph(document, str(summary["takeaway"]), before=0, after=5, size=10.5, color=MUTED)
    takeaway.paragraph_format.keep_with_next = True


def add_problem_method_columns(document: Document, summary: Mapping[str, object]) -> None:
    """Add concise prose sections; tables remain reserved for quantified results."""
    _add_paragraph(document, "Problem and research question", style="Heading 1")
    _add_paragraph(document, str(summary["problem"]), after=3)
    _add_paragraph(document, "Method", style="Heading 1")
    _add_paragraph(document, str(summary["method"]), after=3)


def add_result_strip(document: Document, results: object) -> None:
    if not isinstance(results, list) or len(results) != 3:
        raise SummaryBuildError("summary requires exactly three result records")
    records = [record for record in results if isinstance(record, Mapping)]
    if len(records) != 3 or not all(isinstance(record.get("value"), str) and isinstance(record.get("label"), str) for record in records):
        raise SummaryBuildError("summary result records require value and label strings")
    _add_paragraph(document, "Key results", style="Heading 1", before=4, after=2)
    _add_paragraph(
        document,
        "Evidence snapshot: " + "; ".join(
            f"{record['value']} {record['label']}" for record in records
        ),
        before=0,
        after=2,
        size=8.25,
        color=MUTED,
    )
    table = document.add_table(rows=1, cols=3)
    _set_table_geometry(table, (3504, 3504, 3504))
    row = table.rows[0]
    for cell, record in zip(row.cells, records, strict=True):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade(cell, RESULT_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1)
        value = paragraph.add_run(str(record["value"]))
        _set_run_font(value, 14, bold=True, color=BLUE)
        label = cell.add_paragraph()
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label.paragraph_format.space_before = Pt(0)
        label.paragraph_format.space_after = Pt(0)
        label_run = label.add_run(str(record["label"]))
        _set_run_font(label_run, 8.25, color=MUTED)


def add_figure(document: Document, figure: Path, caption: object) -> None:
    if not figure.is_file():
        raise SummaryBuildError(f"admitted summary figure does not exist: {figure}")
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(4)
    picture.paragraph_format.space_after = Pt(0)
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(str(figure), width=Inches(4.25))
    caption_paragraph = _add_paragraph(document, f"Figure. {caption}", style="Caption", before=1, after=3)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.keep_with_next = True


def add_significance_limitations(document: Document, summary: Mapping[str, object]) -> None:
    _add_paragraph(document, "Why this matters", style="Heading 1", before=2, after=1)
    _add_paragraph(document, str(summary["significance"]), after=2)
    _add_paragraph(document, "Simulation scope", style="Heading 1", before=2, after=1)
    _add_paragraph(document, str(summary["limitations"]), after=2)
    _add_paragraph(document, "Limitations and next steps", style="Heading 1", before=2, after=1)
    _add_paragraph(document, str(summary["nextSteps"]), after=2)


def add_sources(document: Document, sources: object) -> None:
    if not isinstance(sources, list) or not sources or not all(isinstance(source, str) for source in sources):
        raise SummaryBuildError("summary requires a compact source list")
    paragraph = _add_paragraph(document, "Sources: " + " | ".join(sources), before=2, after=0, size=7.5, color=MUTED)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(2)


def _require_summary(package: Mapping[str, object]) -> Mapping[str, object]:
    summary = package.get("summary")
    if not isinstance(summary, Mapping):
        raise SummaryBuildError("Phase 7B package requires a summary object")
    required = {
        "title", "takeaway", "problem", "method", "results", "figure",
        "figureCaption", "significance", "limitations", "nextSteps", "sources",
    }
    if set(summary) != required or not all(isinstance(summary[key], str) for key in required - {"results", "sources"}):
        raise SummaryBuildError("Phase 7B package summary has an invalid shape")
    return summary


def build_summary(root: Path, output: Path) -> None:
    """Create and normalize the admitted one-page DOCX summary."""
    import json
    import tempfile

    root = root.resolve(strict=True)
    package_path = root / "results/presentation/phase7b_package.json"
    package = json.loads(package_path.read_text(encoding="utf-8"))
    if not isinstance(package, Mapping):
        raise SummaryBuildError("Phase 7B package root must be an object")
    summary = _require_summary(package)
    if "{{" in "\n".join(str(value) for value in summary.values()):
        raise SummaryBuildError("Phase 7B package must resolve all content tokens")
    document = Document()
    configure_document(document)
    add_masthead(document, summary)
    add_problem_method_columns(document, summary)
    add_result_strip(document, summary["results"])
    add_figure(document, root / str(summary["figure"]), summary["figureCaption"])
    add_significance_limitations(document, summary)
    add_sources(document, summary["sources"])
    if len(document.sections) != 1 or document.sections[0].start_type != WD_SECTION.NEW_PAGE:
        raise SummaryBuildError("summary must retain one portrait section")
    output = output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(suffix=".docx", dir=output.parent, delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        document.save(temporary_path)
        normalize_openxml_package(temporary_path, ".docx")
        temporary_path.replace(output)
    except BaseException:
        temporary_path.unlink(missing_ok=True)
        raise
