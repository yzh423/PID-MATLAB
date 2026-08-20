"""Deterministic DOCX assembly for the controlled Phase 7A report Markdown."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


PAGE_WIDTH = Inches(8.27)
PAGE_HEIGHT = Inches(11.69)
MARGIN = Inches(0.72)
BODY_FONT = "Aptos"
BODY_SIZE = Pt(9.5)
CAPTION_SIZE = Pt(8.0)
AVAILABLE_WIDTH_INCHES = 8.27 - 2 * 0.72

IMAGE_PATTERN = re.compile(r"^!\[([^]]+)]\(([^)]+)\)$")
HEADING_PATTERN = re.compile(r"^(#{1,3})\s+(.+)$")
ORDERED_ITEM_PATTERN = re.compile(r"^\d+\.\s+(.+)$")
REFERENCE_PATTERN = re.compile(r"^\[(\d+)]\s+(.+)$")
INLINE_PATTERN = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")

TABLE_CAPTIONS = (
    "Frozen robot and actuator parameters.",
    "Controller gain sets and design roles.",
    "Frozen evidence layers and their primary decisions.",
    "Nominal tracking and actuator-demand results.",
    "Constrained PID optimization outcomes.",
    "Deterministic robustness outcomes.",
    "High-noise stochastic robustness outcomes.",
    "Cartesian task acceptance outcomes.",
    "Cross-model validation summary.",
)


class DocumentBuildError(ValueError):
    """Raised when controlled Markdown cannot be assembled safely."""


@dataclass(frozen=True)
class DocumentMetadata:
    figure_paths: tuple[str, ...]
    figure_count: int
    table_count: int


def build_docx(
    markdown: str, project_root: Path, output_path: Path
) -> DocumentMetadata:
    """Build and atomically replace one styled report DOCX."""
    project_root = project_root.resolve(strict=True)
    output_path = output_path.resolve()
    if output_path.suffix.lower() != ".docx":
        raise DocumentBuildError("DOCX output must use the .docx extension")
    if not _is_within(output_path, project_root / "docs/report"):
        raise DocumentBuildError("DOCX output must stay inside docs/report")

    document = Document()
    _configure_document(document)
    metadata = _assemble_markdown(
        document, markdown, project_root, project_root / "docs/report"
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        suffix=".docx", dir=output_path.parent, delete=False
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        document.save(temporary_path)
        _normalize_docx_archive(temporary_path)
        Document(temporary_path)
        temporary_path.replace(output_path)
    except BaseException:
        temporary_path.unlink(missing_ok=True)
        raise
    return metadata


def _configure_document(document: DocumentType) -> None:
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    section.different_first_page_header_footer = True
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = RGBColor(31, 41, 55)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    _set_style(styles["Title"], 19, bold=True, color="17365D", after=8)
    _set_style(styles["Subtitle"], 10.5, color="4B5563", after=4)
    _set_style(styles["Heading 1"], 15, bold=True, color="17365D", before=10, after=4)
    _set_style(styles["Heading 2"], 12, bold=True, color="1F4E79", before=8, after=3)
    _set_style(styles["Heading 3"], 10.5, bold=True, color="365F91", before=6, after=2)
    _set_style(styles["Caption"], 8, italic=True, color="4B5563", before=2, after=5)
    styles["Caption"].paragraph_format.keep_with_next = False

    if "Report Metadata" not in styles:
        metadata = styles.add_style("Report Metadata", WD_STYLE_TYPE.PARAGRAPH)
    else:
        metadata = styles["Report Metadata"]
    _set_style(metadata, 9.5, color="4B5563", after=2)

    if "Protocol Block" not in styles:
        protocol = styles.add_style("Protocol Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        protocol = styles["Protocol Block"]
    protocol.font.name = "Aptos Mono"
    protocol.font.size = Pt(8.5)
    protocol.font.color.rgb = RGBColor(31, 41, 55)
    protocol.paragraph_format.left_indent = Inches(0.2)
    protocol.paragraph_format.right_indent = Inches(0.2)
    protocol.paragraph_format.space_before = Pt(3)
    protocol.paragraph_format.space_after = Pt(6)

    document.core_properties.title = (
        "Reliable Robotic Manipulation Through Evidence-Grounded PID and Fuzzy-PID Evaluation"
    )
    document.core_properties.subject = "Phase 7A technical report"
    document.core_properties.author = "PID vs Fuzzy PID project"
    _configure_header_footer(section)


def _set_style(
    style: object,
    size: float,
    *,
    bold: bool = False,
    italic: bool = False,
    color: str = "1F2937",
    before: float = 0,
    after: float = 0,
) -> None:
    style.font.name = BODY_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True


def _configure_header_footer(section: object) -> None:
    _set_header(section.header.paragraphs[0])
    _set_header(section.even_page_header.paragraphs[0])
    section.first_page_header.paragraphs[0].clear()
    _set_footer(section.footer.paragraphs[0])
    _set_footer(section.even_page_footer.paragraphs[0])
    _set_footer(section.first_page_footer.paragraphs[0])


def _set_header(header: object) -> None:
    header.text = "PID vs Fuzzy PID  |  Phase 7A Technical Report"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.runs[0]
    header_run.font.name = BODY_FONT
    header_run.font.size = Pt(7.5)
    header_run.font.color.rgb = RGBColor(107, 114, 128)



def _set_footer(footer: object) -> None:
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Page ")
    run.font.name = BODY_FONT
    run.font.size = Pt(8)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend((field_begin, instruction, field_end))


def _assemble_markdown(
    document: DocumentType, markdown: str, project_root: Path, report_dir: Path
) -> DocumentMetadata:
    lines = markdown.splitlines()
    index = 0
    figure_paths: list[str] = []
    figure_number = 0
    table_number = 0
    in_references = False

    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue

        heading = HEADING_PATTERN.match(line)
        if heading:
            level = len(heading.group(1))
            text = _plain_text(heading.group(2))
            if level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.paragraph_format.space_before = Pt(34)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.add_run(text)
            else:
                if text == "References":
                    in_references = True
                paragraph = document.add_paragraph(text, style=f"Heading {level - 1}")
            index += 1
            continue

        if line.startswith("```"):
            if line != "```text":
                raise DocumentBuildError(f"unsupported fenced block: {line}")
            index += 1
            block: list[str] = []
            while index < len(lines) and lines[index].rstrip() != "```":
                block.append(lines[index].rstrip())
                index += 1
            if index >= len(lines):
                raise DocumentBuildError("unterminated text fence")
            paragraph = document.add_paragraph(style="Protocol Block")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run("\n".join(block))
            _shade_paragraph(paragraph, "F3F6FA")
            index += 1
            continue

        image = IMAGE_PATTERN.match(line)
        if image:
            caption, relative_path = image.groups()
            image_path = (report_dir / relative_path).resolve()
            if not _is_within(image_path, project_root):
                raise DocumentBuildError(f"image escapes project root: {relative_path}")
            if not image_path.is_file() or image_path.suffix.lower() != ".png":
                raise DocumentBuildError(f"missing PNG image: {relative_path}")
            figure_number += 1
            image_paragraph = document.add_paragraph()
            image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image_paragraph.paragraph_format.keep_with_next = True
            with Image.open(image_path) as source_image:
                aspect_ratio = source_image.width / source_image.height
            width_inches = min(6.45, 2.95 * aspect_ratio)
            image_paragraph.add_run().add_picture(
                str(image_path), width=Inches(width_inches)
            )
            caption_paragraph = document.add_paragraph(style="Caption")
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption_paragraph.add_run(f"Figure {figure_number}. {caption}")
            figure_paths.append(image_path.relative_to(project_root).as_posix())
            index += 1
            continue

        if _looks_like_table(lines, index):
            table_rows, index = _consume_table(lines, index)
            table_number += 1
            if table_number > len(TABLE_CAPTIONS):
                raise DocumentBuildError("table caption map is incomplete")
            caption = document.add_paragraph(style="Caption")
            caption.paragraph_format.keep_with_next = True
            caption.add_run(f"Table {table_number}. {TABLE_CAPTIONS[table_number - 1]}")
            _add_table(document, table_rows)
            continue

        if line.startswith(">"):
            paragraph = document.add_paragraph(style="Intense Quote")
            _add_inline(paragraph, line[1:].strip())
            index += 1
            continue

        if line.startswith(("- ", "* ")):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline(paragraph, line[2:].strip())
            index += 1
            continue

        ordered = ORDERED_ITEM_PATTERN.match(line)
        if ordered:
            paragraph = document.add_paragraph(style="List Number")
            _add_inline(paragraph, ordered.group(1))
            index += 1
            continue

        reference = REFERENCE_PATTERN.match(line) if in_references else None
        if reference:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_after = Pt(3)
            _add_inline(paragraph, line)
            index += 1
            continue

        if line.startswith("**") and "**" in line[2:]:
            paragraph = document.add_paragraph(style="Report Metadata")
            _add_inline(paragraph, line)
            index += 1
            continue

        if line.startswith("<") or re.search(r"</?[A-Za-z][^>]*>", line):
            raise DocumentBuildError("HTML is not supported in report Markdown")

        paragraph_lines = [line]
        index += 1
        while index < len(lines) and lines[index].strip():
            candidate = lines[index].rstrip()
            if (
                HEADING_PATTERN.match(candidate)
                or candidate.startswith("```")
                or IMAGE_PATTERN.match(candidate)
                or _looks_like_table(lines, index)
                or candidate.startswith((">", "- ", "* "))
                or ORDERED_ITEM_PATTERN.match(candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.widow_control = True
        _add_inline(paragraph, " ".join(paragraph_lines))

    if figure_number != 6:
        raise DocumentBuildError(f"expected six figures, found {figure_number}")
    if table_number != len(TABLE_CAPTIONS):
        raise DocumentBuildError(
            f"expected {len(TABLE_CAPTIONS)} tables, found {table_number}"
        )
    return DocumentMetadata(tuple(figure_paths), figure_number, table_number)


def _looks_like_table(lines: list[str], index: int) -> bool:
    return (
        index + 1 < len(lines)
        and lines[index].lstrip().startswith("|")
        and lines[index + 1].lstrip().startswith("|")
        and all(
            re.fullmatch(r":?-{3,}:?", cell.strip())
            for cell in _split_table_row(lines[index + 1])
        )
    )


def _consume_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    header = _split_table_row(lines[index])
    separator = _split_table_row(lines[index + 1])
    if len(header) != len(separator):
        raise DocumentBuildError("malformed Markdown table width")
    rows = [header]
    index += 2
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        row = _split_table_row(lines[index])
        if len(row) != len(header):
            raise DocumentBuildError("malformed Markdown table width")
        rows.append(row)
        index += 1
    return rows, index


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _add_table(document: DocumentType, rows: list[list[str]]) -> None:
    column_count = len(rows[0])
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    font_size = 7.3 if column_count >= 6 else 8.0 if column_count >= 5 else 8.5
    width = Inches(AVAILABLE_WIDTH_INCHES / column_count)

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        _prevent_row_split(row)
        if row_index == 0:
            _repeat_table_header(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, top=45, start=55, bottom=45, end=55)
            if row_index == 0:
                _shade_cell(cell, "E8EEF5")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.keep_with_next = row_index == 0
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            _add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = BODY_FONT
                run.font.size = Pt(font_size)
                if row_index == 0:
                    run.bold = True
    spacer = document.add_paragraph()
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    spacer.paragraph_format.line_spacing = Pt(2.5)
    spacer.paragraph_format.space_after = Pt(0)


def _add_inline(paragraph: object, text: str) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Aptos Mono"
            run.font.size = Pt(8.5)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def _plain_text(text: str) -> str:
    return re.sub(r"(\*\*|\*|`)", "", text).strip()


def _shade_paragraph(paragraph: object, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _shade_cell(cell: object, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _repeat_table_header(row: object) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _prevent_row_split(row: object) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)


def _set_cell_margins(cell: object, **kwargs: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        if edge not in kwargs:
            continue
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(kwargs[edge]))
        node.set(qn("w:type"), "dxa")


def _is_within(path: Path, directory: Path) -> bool:
    try:
        path.resolve().relative_to(directory.resolve())
    except ValueError:
        return False
    return True


def _normalize_docx_archive(path: Path) -> None:
    """Rewrite the OOXML package with stable ordering and ZIP timestamps."""
    normalized_path = path.with_name(path.stem + ".normalized.docx")
    try:
        with ZipFile(path, "r") as source:
            entries = [(item, source.read(item.filename)) for item in source.infolist()]
        with ZipFile(normalized_path, "w", compression=ZIP_DEFLATED, compresslevel=9) as target:
            for source_info, payload in sorted(entries, key=lambda item: item[0].filename):
                info = ZipInfo(source_info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                info.compress_type = ZIP_DEFLATED
                info.create_system = 0
                info.external_attr = source_info.external_attr
                target.writestr(info, payload)
        normalized_path.replace(path)
    except BaseException:
        normalized_path.unlink(missing_ok=True)
        raise
