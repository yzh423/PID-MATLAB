from __future__ import annotations

import hashlib
import json
from pathlib import Path
import re
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from PIL import Image, PngImagePlugin
from pypdf import PdfReader

from scripts import build_report as report_builder
from scripts.reporting import document as report_document
from scripts.reporting.document import build_docx
from scripts.normalize_report_pdf import FIXED_PDF_DATE, normalize_pdf


PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = PROJECT_ROOT / "docs/report"
DOCX_PATH = REPORT_DIR / "technical_report.docx"
PDF_PATH = REPORT_DIR / "technical_report.pdf"
MANIFEST_PATH = REPORT_DIR / "build_manifest.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class BuiltDocumentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if not DOCX_PATH.is_file():
            raise AssertionError("technical_report.docx not built")
        cls.document = Document(DOCX_PATH)
        cls.text = "\n".join(paragraph.text for paragraph in cls.document.paragraphs)

    def test_page_geometry_and_typography(self) -> None:
        self.assertGreaterEqual(len(self.document.sections), 1)
        for section in self.document.sections:
            self.assertAlmostEqual(section.page_width.inches, 8.27, places=2)
            self.assertAlmostEqual(section.page_height.inches, 11.69, places=2)
            for margin in (
                section.top_margin,
                section.bottom_margin,
                section.left_margin,
                section.right_margin,
            ):
                self.assertGreaterEqual(margin.inches, 0.65)
                self.assertLessEqual(margin.inches, 0.85)

        normal_size = self.document.styles["Normal"].font.size
        caption_size = self.document.styles["Caption"].font.size
        self.assertIsNotNone(normal_size)
        self.assertIsNotNone(caption_size)
        self.assertGreaterEqual(normal_size.pt, 9.5)
        self.assertGreaterEqual(caption_size.pt, 8.0)

    def test_required_structure_and_metadata(self) -> None:
        self.assertIn("Reliable Robotic Manipulation", self.text)
        self.assertIn("Technical report — Phase 7A", self.text)
        self.assertIn("PID vs Fuzzy PID for a two-link robotic manipulator", self.text)
        for heading in (
            "Abstract",
            "Introduction and Research Context",
            "Problem Formulation",
            "System Model",
            "Controller Design and Tuning",
            "Experimental Design",
            "Results",
            "Independent Model Validation",
            "Discussion",
            "Embodied-AI Execution Context",
            "Limitations and Future Work",
            "Conclusion",
            "References",
        ):
            self.assertIn(heading, self.text)
        self.assertNotIn("{{", self.text)
        self.assertNotIn("}}", self.text)

    def test_images_tables_and_captions_are_complete(self) -> None:
        self.assertGreaterEqual(len(self.document.inline_shapes), 6)
        self.assertGreaterEqual(len(self.document.tables), 7)

        figure_numbers = []
        table_numbers = []
        for paragraph in self.document.paragraphs:
            if match := re.match(r"Figure (\d+)\.", paragraph.text):
                figure_numbers.append(int(match.group(1)))
            if match := re.match(r"Table (\d+)\.", paragraph.text):
                table_numbers.append(int(match.group(1)))
        self.assertEqual(figure_numbers, list(range(1, len(figure_numbers) + 1)))
        self.assertEqual(table_numbers, list(range(1, len(table_numbers) + 1)))
        self.assertEqual(len(figure_numbers), 6)
        self.assertGreaterEqual(len(table_numbers), 7)

        with ZipFile(DOCX_PATH) as archive:
            relationships = archive.read("word/_rels/document.xml.rels").decode("utf-8")
        image_relationships = re.findall(r'<Relationship[^>]+Type="[^"]+/image"[^>]*/>', relationships)
        self.assertGreaterEqual(len(image_relationships), 6)
        self.assertTrue(all('TargetMode="External"' not in item for item in image_relationships))

    def test_references_are_numbered_and_complete(self) -> None:
        for citation_id in range(1, 9):
            self.assertRegex(self.text, rf"(?m)^\[{citation_id}\] ")

    def test_manifest_hashes_current_inputs_and_outputs(self) -> None:
        self.assertTrue(MANIFEST_PATH.is_file(), "build_manifest.json not built")
        manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
        self.assertEqual(manifest["schemaVersion"], 1)
        self.assertEqual(len(manifest["selectedFigures"]), 6)
        self.assertGreater(len(manifest["usedEvidenceTokens"]), 0)
        for group in ("sources", "outputs"):
            self.assertGreater(len(manifest[group]), 0)
            for record in manifest[group]:
                path = PROJECT_ROOT / record["path"]
                self.assertTrue(path.is_file(), record["path"])
                self.assertEqual(record["sha256"], sha256(path), record["path"])

    def test_final_pdf_is_complete_and_extractable(self) -> None:
        self.assertTrue(PDF_PATH.is_file(), "technical_report.pdf not exported")
        reader = PdfReader(PDF_PATH)
        self.assertGreaterEqual(len(reader.pages), 8)
        self.assertLessEqual(len(reader.pages), 12)
        page_text = [(page.extract_text() or "").strip() for page in reader.pages]
        self.assertTrue(all(len(text) > 40 for text in page_text))
        text = "\n".join(page_text)
        self.assertIn("Reliable Robotic Manipulation", text)
        self.assertIn("Limitations and Future Work", text)
        self.assertIn("References", text)
        self.assertNotIn("{{", text)
        self.assertNotIn("}}", text)
        manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
        self.assertIn(
            "docs/report/technical_report.pdf",
            [record["path"] for record in manifest["outputs"]],
        )
        for citation_id in range(1, 9):
            self.assertRegex(text, rf"(?m)^\[{citation_id}\] ")

        self.assertEqual(reader.metadata.creation_date.strftime("D:%Y%m%d%H%M%S%z"), "D:20260820000000+0800")

    def test_docx_and_pdf_packaging_is_binary_deterministic(self) -> None:
        markdown = (REPORT_DIR / "technical_report.md").read_text(encoding="utf-8")
        with tempfile.TemporaryDirectory(dir=REPORT_DIR) as temporary_directory:
            temporary = Path(temporary_directory)
            first_docx = temporary / "first.docx"
            second_docx = temporary / "second.docx"
            build_docx(markdown, PROJECT_ROOT, first_docx)
            build_docx(markdown, PROJECT_ROOT, second_docx)
            self.assertEqual(sha256(first_docx), sha256(second_docx))

            first_pdf = temporary / "first.pdf"
            second_pdf = temporary / "second.pdf"
            normalize_pdf(PDF_PATH, first_pdf, PROJECT_ROOT)
            normalize_pdf(PDF_PATH, second_pdf, PROJECT_ROOT)
            self.assertEqual(sha256(first_pdf), sha256(second_pdf))
            self.assertEqual(PdfReader(first_pdf).metadata.get("/CreationDate"), FIXED_PDF_DATE)

    def test_png_normalization_removes_volatile_matlab_metadata(self) -> None:
        with tempfile.TemporaryDirectory(dir=REPORT_DIR) as temporary_directory:
            temporary = Path(temporary_directory)
            paths = [temporary / "first.png", temporary / "second.png"]
            for path, timestamp in zip(paths, ("first-run", "second-run"), strict=True):
                metadata = PngImagePlugin.PngInfo()
                metadata.add_text("Software", "MATLAB, The MathWorks, Inc.")
                metadata.add_text("Creation Time", timestamp)
                Image.new("RGB", (4, 3), (12, 34, 56)).save(
                    path, format="PNG", pnginfo=metadata, dpi=(180, 180)
                )

            self.assertNotEqual(sha256(paths[0]), sha256(paths[1]))
            for path in paths:
                report_document._normalize_png_file(path)

            self.assertEqual(sha256(paths[0]), sha256(paths[1]))
            for path in paths:
                with Image.open(path) as normalized:
                    self.assertNotIn("Creation Time", normalized.info)

    def test_phase7a_text_files_pin_lf_checkout_endings(self) -> None:
        attributes_path = PROJECT_ROOT / ".gitattributes"
        self.assertTrue(attributes_path.is_file())
        attributes = attributes_path.read_text(encoding="utf-8")
        for pattern in (
            "/+rrm/+report/** text eol=lf",
            "/docs/report/*.md text eol=lf",
            "/docs/report/*.json text eol=lf",
            "/docs/report/*.docx binary",
            "/docs/report/*.pdf binary",
            "/scripts/reporting/** text eol=lf",
            "/tests/report/** text eol=lf",
        ):
            self.assertIn(pattern, attributes)

    def test_report_source_normalization_rewrites_crlf_bytes(self) -> None:
        with tempfile.TemporaryDirectory(dir=REPORT_DIR) as temporary_directory:
            source = Path(temporary_directory) / "source.md"
            source.write_bytes(b"first\r\nsecond\r\n")

            report_builder._normalize_text_file(source)

            self.assertEqual(source.read_bytes(), b"first\nsecond\n")


if __name__ == "__main__":
    unittest.main()
