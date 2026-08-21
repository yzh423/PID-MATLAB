from __future__ import annotations

from pathlib import Path
import re
import subprocess
import tempfile
import unittest
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
PAGINATOR = ROOT / "scripts/phase7b/measure_docx_pages.ps1"


def word_process_ids() -> set[int]:
    completed = subprocess.run(
        [
            "powershell", "-NoProfile", "-Command",
            "@(Get-Process -Name WINWORD -ErrorAction SilentlyContinue | Select-Object -ExpandProperty Id) -join ','",
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    return {int(value) for value in completed.stdout.strip().split(",") if value}


class Phase7BWordPaginationTests(unittest.TestCase):
    def test_three_rendered_pages_fail_the_one_page_contract_even_when_metadata_says_one(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            docx = Path(directory) / "metadata-lies.docx"
            document = Document()
            document.add_paragraph("Page one")
            document.add_page_break()
            document.add_paragraph("Page two")
            document.add_page_break()
            document.add_paragraph("Page three")
            document.save(docx)
            with ZipFile(docx) as archive:
                properties = ElementTree.fromstring(archive.read("docProps/app.xml"))
            pages = next(element for element in properties.iter() if element.tag.endswith("}Pages"))
            self.assertEqual(pages.text, "1", "fixture must reproduce stale OOXML metadata")

            before = word_process_ids()
            completed = subprocess.run(
                [
                    "powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
                    "-File", str(PAGINATOR), "-DocxPath", str(docx), "-TimeoutSeconds", "30",
                ],
                cwd=ROOT,
                check=False,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=45,
            )
            after = word_process_ids()
            self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
            match = re.search(r"DOCX_RENDERED_PAGE_COUNT=(\d+)", completed.stdout)
            self.assertIsNotNone(match, completed.stdout)
            assert match is not None
            self.assertEqual(int(match.group(1)), 3)
            self.assertEqual(after, before, "pagination must not leak or terminate WINWORD processes")

    def test_paginator_is_pid_scoped_hidden_read_only_and_fail_closed(self) -> None:
        self.assertTrue(PAGINATOR.is_file(), "rendered pagination worker is required")
        source = PAGINATOR.read_text(encoding="utf-8")
        for token in (
            "$word.Visible = $false",
            "OpenNoRepairDialog($resolvedDocx, $false, $true, $false)",
            "$document.Repaginate()",
            "$document.ComputeStatistics(2)",
            "ownedWordPid",
            "preexistingWordPids",
            "WaitForExit",
            "TimeoutSeconds",
        ):
            self.assertIn(token, source)
        self.assertNotIn("Stop-Process -Name WINWORD", source)


if __name__ == "__main__":
    unittest.main()
