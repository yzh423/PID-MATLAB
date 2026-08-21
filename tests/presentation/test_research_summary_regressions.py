from __future__ import annotations

import json
import re
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document
import pdfplumber
from pypdf import PdfReader

from scripts.phase7b.evidence import sha256_file


ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "docs" / "summary" / "research_summary.docx"
PDF = ROOT / "docs" / "summary" / "research_summary.pdf"
EXPORTER = ROOT / "scripts" / "export_research_summary_pdf.ps1"
DOCX_BUILDER = ROOT / "scripts" / "build_research_summary.py"
PDF_BUILDER = ROOT / "scripts" / "build_research_summary_pdf.py"
PACKAGE = ROOT / "results" / "presentation" / "phase7b_package.json"
DESIGN = ROOT / "docs" / "skills" / "specs" / "2026-08-21-phase-7b-presentation-summary-design.md"
PLAN = ROOT / "docs" / "skills" / "plans" / "2026-08-21-phase-7b-presentation-summary.md"
FIGURE_RATIO = 1825 / 1171
NAMESPACES = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


class ResearchSummaryOOXMLTests(unittest.TestCase):
    def test_docx_geometry_table_figure_and_text_contract(self) -> None:
        document = Document(DOCX)
        self.assertEqual(len(document.sections), 1)
        section = document.sections[0]
        self.assertAlmostEqual(section.page_width.inches, 8.5, places=2)
        self.assertAlmostEqual(section.page_height.inches, 11, places=2)
        self.assertAlmostEqual(section.left_margin.inches, 0.6, places=2)
        self.assertAlmostEqual(section.right_margin.inches, 0.6, places=2)
        self.assertAlmostEqual(section.top_margin.inches, 0.55, places=2)
        self.assertAlmostEqual(section.bottom_margin.inches, 0.55, places=2)
        self.assertEqual(len(document.inline_shapes), 1)
        shape = document.inline_shapes[0]
        self.assertAlmostEqual(shape.width / shape.height, FIGURE_RATIO, places=2)
        with ZipFile(DOCX) as archive:
            names = archive.namelist()
            payload = archive.read("word/document.xml")
        self.assertFalse(any(name.startswith("word/comments") for name in names))
        self.assertNotIn(b"w:ins", payload)
        self.assertNotIn(b"w:del", payload)
        self.assertNotIn(b"{{", payload)
        root = ElementTree.fromstring(payload)
        table = root.find(".//w:tbl", NAMESPACES)
        self.assertIsNotNone(table)
        assert table is not None
        table_properties = table.find("./w:tblPr", NAMESPACES)
        self.assertEqual(table_properties.find("./w:tblW", NAMESPACES).get(f"{{{NAMESPACES['w']}}}w"), "10512")
        self.assertEqual(table_properties.find("./w:tblInd", NAMESPACES).get(f"{{{NAMESPACES['w']}}}w"), "0")
        self.assertEqual(
            [column.get(f"{{{NAMESPACES['w']}}}w") for column in table.findall("./w:tblGrid/w:gridCol", NAMESPACES)],
            ["3504", "3504", "3504"],
        )
        self.assertEqual(
            [cell.get(f"{{{NAMESPACES['w']}}}w") for cell in table.findall(".//w:tcPr/w:tcW", NAMESPACES)],
            ["3504", "3504", "3504"],
        )


class ResearchSummaryPDFLayoutTests(unittest.TestCase):
    def test_pdf_preserves_image_aspect_embeds_font_and_uses_page_depth(self) -> None:
        reader = PdfReader(PDF)
        self.assertEqual(len(reader.pages), 1)
        page = reader.pages[0]
        media_box = page.mediabox
        self.assertAlmostEqual(float(media_box.width), 612, places=1)
        self.assertAlmostEqual(float(media_box.height), 792, places=1)
        contents = page.get_contents().get_data().decode("latin-1")
        image_matrix = re.search(
            r"q\s+([\d.]+)\s+0\s+0\s+([\d.]+)\s+[\d.]+\s+[\d.]+\s+cm\s+/\S+\s+Do",
            contents,
        )
        self.assertIsNotNone(image_matrix, "PDF must draw the admitted figure as an image XObject")
        assert image_matrix is not None
        self.assertAlmostEqual(float(image_matrix.group(1)) / float(image_matrix.group(2)), FIGURE_RATIO, delta=0.02)
        self.assertIn("10.25 Tf", contents, "fallback body typography must match the DOCX 10.25 pt body")
        fonts = page["/Resources"]["/Font"].get_object().values()
        self.assertTrue(
            any("/FontDescriptor" in font.get_object() for font in fonts),
            "fallback must embed an Arial-compatible font rather than rely on base-14 Helvetica",
        )
        text = page.extract_text() or ""
        normalized_text = re.sub(r"\s+", " ", text)
        for claim in (
            "7.283%",
            "10/13",
            "30/30",
            "0/30",
            "optimized PID deterministic scenarios passed",
            "per controller-scenario cell: isolated-noise cell vs full-combined-stress cell",
            "docs/report/build_manifest.json",
            "Simulation scope",
            "Limitations and next steps",
        ):
            self.assertIn(claim, normalized_text)
        with pdfplumber.open(PDF) as document:
            source_words = [word for word in document.pages[0].extract_words() if word["text"].startswith("Sources:")]
        self.assertEqual(len(source_words), 1)
        self.assertGreater(source_words[0]["bottom"], 650, "content should use the lower Letter page without a top-heavy blank third")
        self.assertLess(source_words[0]["bottom"], 752, "source footer must remain inside the bottom margin")

    def test_docx_and_pdf_are_semantically_and_structurally_equivalent_renderers(self) -> None:
        package = json.loads(PACKAGE.read_text(encoding="utf-8"))
        summary = package["summary"]
        document = Document(DOCX)
        docx_text = " ".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        pdf_text = " ".join((page.extract_text() or "") for page in PdfReader(PDF).pages)

        def normalized(value: str) -> str:
            return re.sub(r"\s+", " ", value).strip()

        for field in ("title", "takeaway", "problem", "method", "significance", "limitations", "nextSteps"):
            expected = normalized(str(summary[field]))
            self.assertIn(expected, normalized(docx_text), field)
            self.assertIn(expected, normalized(pdf_text), field)
        for result in summary["results"]:
            for value in (result["value"], result["label"]):
                self.assertIn(normalized(value), normalized(docx_text))
                self.assertIn(normalized(value), normalized(pdf_text))
        for source in summary["sources"]:
            self.assertIn(source, docx_text)
            self.assertIn(source, pdf_text)

        docx_order = [docx_text.index(text) for text in (
            "Problem and research question", "Method", "Key results", "Why this matters", "Simulation scope", "Limitations and next steps", "Sources:"
        )]
        pdf_order = [pdf_text.index(text) for text in (
            "Problem and research question", "Method", "Key results", "Why this matters", "Simulation scope", "Limitations and next steps", "Sources:"
        )]
        self.assertEqual(docx_order, sorted(docx_order))
        self.assertEqual(pdf_order, sorted(pdf_order))

    def test_contract_defines_two_independent_canonical_renderers(self) -> None:
        design = DESIGN.read_text(encoding="utf-8")
        plan = PLAN.read_text(encoding="utf-8")
        for source in (design, plan):
            self.assertIn("two independent canonical renderers", source)
            self.assertIn("same deterministic Phase 7B package", source)
            self.assertNotIn("PDF export: owned hidden Microsoft Word COM process", source)
        self.assertIn("results/presentation/phase7b_package.json", PDF_BUILDER.read_text(encoding="utf-8"))
        self.assertIn("results/presentation/phase7b_package.json", (ROOT / "scripts/phase7b/summary.py").read_text(encoding="utf-8"))

    def test_plan_and_design_describe_the_current_keyed_audit_gated_contract(self) -> None:
        design = DESIGN.read_text(encoding="utf-8")
        plan = PLAN.read_text(encoding="utf-8")
        self.assertIn("Optimized PID is the strongest reliability baseline", design)
        self.assertIn("actual rendered pagination", design)
        self.assertIn("OOXML-resolved geometry", design)
        self.assertNotIn("cartesian.runRows.5", plan)
        self.assertIn("phase7b.cartesian.optimizedPickTransferPlace", plan)
        self.assertNotIn("docs/report/PAPER_CLAIM_AUDIT.md", plan)
        self.assertNotIn("validate bundled Node/Python and Word availability", plan)
        self.assertNotIn("export PDF, and normalize both", plan)
        self.assertIn("canonical Phase 7B audit gate", plan)


class ResearchSummaryExportLifecycleTests(unittest.TestCase):
    def test_exporter_has_independent_com_and_temp_cleanup_guarantees(self) -> None:
        source = EXPORTER.read_text(encoding="utf-8")
        for function in (
            "function Close-WordDocument",
            "function Quit-WordApplication",
            "function Release-ComReference",
            "function Remove-TaskOwnedTemporaryPdf",
        ):
            self.assertIn(function, source)
        self.assertRegex(source, r"try \{ Close-WordDocument .*? \} catch \{ \[void\]\$cleanupErrors\.Add")
        self.assertRegex(source, r"try \{ Quit-WordApplication .*? \} catch \{ \[void\]\$cleanupErrors\.Add")
        self.assertIn("Word diagnostic cleanup failed closed", source)
        self.assertRegex(
            source,
            r"finally \{\s*Remove-TaskOwnedTemporaryPdf -Path \$temporaryPdf -SummaryRoot \$summaryRoot",
        )
        self.assertIn(".research_summary.", source)
        self.assertIn(".tmp.pdf", source)
        self.assertIn("non-canonical diagnostic", source)
        self.assertIn("ownedProcessIds", source)
        self.assertIn("Owned WINWORD process remains", source)

    def test_word_diagnostic_requires_explicit_noncanonical_output(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            summary = root / "docs/summary"
            summary.mkdir(parents=True)
            shutil.copy2(DOCX, summary / DOCX.name)
            shutil.copy2(PDF, summary / PDF.name)
            canonical = summary / PDF.name
            original_hash = sha256_file(canonical)
            cases = (
                ["-UseWordCom"],
                ["-UseWordCom", "-PdfPath", str(canonical)],
            )
            for arguments in cases:
                with self.subTest(arguments=arguments):
                    completed = subprocess.run(
                        [
                            "powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
                            "-File", str(EXPORTER), "-ProjectRoot", str(root), *arguments,
                        ],
                        cwd=ROOT,
                        check=False,
                        capture_output=True,
                        text=True,
                        timeout=45,
                    )
                    self.assertNotEqual(completed.returncode, 0, completed.stdout)
                    self.assertRegex(
                        completed.stdout + completed.stderr,
                        "explicit.*diagnostic|canonical",
                    )
                    self.assertEqual(sha256_file(canonical), original_hash)


if __name__ == "__main__":
    unittest.main()
